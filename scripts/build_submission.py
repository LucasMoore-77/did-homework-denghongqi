"""Build print submission from recorded evidence; uses python-docx."""
from pathlib import Path
import json,subprocess,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'submission';OUT.mkdir(exist_ok=True)
meta=json.loads((ROOT/'config/submission.json').read_text());doc=Document();sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Noto Serif CJK SC';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(11 if name=='Normal' else 23 if name=='Title' else 16 if name=='Heading 1' else 12)
 st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=Pt(16 if name=='Normal' else 28 if name=='Title' else 21 if name=='Heading 1' else 17)
 if name.startswith('Heading'):st.font.bold=True
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=foot.add_run('邓鸿琪  825200520  ·  ');run.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
md=[]
pending_break=False
def title(s):doc.add_paragraph(s,'Title');md.append('# '+s+'\n')
def h(s,level=1):
 global pending_break
 pp=doc.add_heading(s,level)
 if pending_break:pp.paragraph_format.page_break_before=True;pending_break=False
 md.append('#'*(level+1)+' '+s+'\n')
def p(s,small=False):
 pp=doc.add_paragraph(s);md.append(s+'\n')
 if small:
  for x in pp.runs:x.font.size=Pt(9)
 return pp
def table(head,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(head));t.style='Table Grid';t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,x in enumerate(head):t.rows[0].cells[i].text=str(x)
 trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 for cell in t.rows[0].cells:
  shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E9EEF2');cell._tc.get_or_add_tcPr().append(shd)
  for rr in cell.paragraphs[0].runs:rr.bold=True
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=str(v)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.line_spacing=Pt(14)
    for rr in pp.runs:rr.font.size=Pt(10)
 md.append('| '+' | '.join(head)+' |\n| '+' | '.join(['---']*len(head))+' |\n'+'\n'.join('| '+' | '.join(str(x).replace('\n','<br>') for x in row)+' |' for row in rows)+'\n')
 return t
def page():
 global pending_break
 pending_break=True;md.append('\n<!-- pagebreak -->\n')
def prompt(code):
 s=(ROOT/f'prompts/robustness/{code}.md').read_text();s=s.split('## Prompt（执行前保存的任务指令）')[1].split('## 结果摘要与观察')[0].strip()
 for para in s.split('\n\n'):p(para)

title('AI时代经济学研究工具与方法')
p('2026年暑期课程期末作业纸质提交材料')
table(['项目','内容'],[['姓名','邓鸿琪'],['学号','825200520'],['学院及专业','江西财经大学数字经济学院\n2025级产业经济学硕士研究生'],['GitHub仓库',meta['github_url']],['提交日期',meta['date']],['用AI完成的工作','复现DID，完成11项必做及1项扩展检验；设计Skill与评估Agent；以两轮对照和错误输入测试完善验证，并形成可编译论文。']],[3.0,13.3])
p('主要发现：基准系数0.120920；短窗口与聚类推断存在敏感性，管理能力互补性未获支持。全部数据为合成教学数据，不构成真实企业因果证据。')
h('一、识别目标')
table(['项目','我的回答'],[['处理与对照','151家企业于2020年采用数字化，209家从未采用；面板共360家、3240个观测。'],['处理时间','2016—2024年观察，2020年同步处理；不存在交错采用。'],['目标估计量','处理组2020—2024年平均ATT。基准系数仅在识别假设成立时与该目标对应；改变时间窗口会改变平均对象。'],['关键假设','条件平行趋势、无预期、无溢出、处理定义一致、样本构成稳定；控制项不为处理后中介。'],['主要威胁','采用自选择、企业时变混杂、小簇推断、能力测量误差。DGP含每年0.0008额外趋势，预趋势不显著不能证明假设成立。']],[3.0,13.3])
p('本材料由Codex辅助分析与草拟，记录以实际文件和执行证据为准；署名研究者提交前复核。',True)
page();h('二、我设计的Prompt')
h('示例A  R2 随机伪处理组安慰剂',2);prompt('R2')
p('设计理由：固定分配单位、种子和样本，保留全部重复，使后续Skill对照可比较。')
h('示例B  T3 能力互补性',2);prompt('T3')
p('设计理由：在运行前规定处理前能力和中心化，并明确交互显著也不等于机制因果识别。')
h('Prompt设计心得',2)
p('最容易跑偏的是把“检验不显著”解释为“假设成立”。R1要求联合检验及区间，R5要求簇数和二维定义，T3要求不显著也如实写出。共用四要素只统一结构，各检验的验证项围绕不同识别威胁设计。R2初版的直接尾部率0在第一轮修订中增加了有限模拟限定，未据此改动回归结果。')
page();h('三、我设计的Skill')
h('3.1 六模块摘要',2)
table(['模块','设计要点'],[['触发条件','面板DID复现与稳健性诊断；不替代其他识别方法。'],['背景知识','先区分ATT、处理暴露、处理组、时点、FE和聚类；记录估计量与假设。'],['工作步骤','读取说明和数据→确认变量语义→生成外部契约→校验时序→复现基线→逐项运行→保存证据。'],['检查清单','唯一键、缺失、处理反转、同步采用、平衡面板、吸收项、处理前分组、少簇、全量重复。'],['边界条件','数据只读，缺失语义先询问；模拟不解释为现实因果证据；交错采用转用合适估计器。'],['验证方式','独立实现对照、哈希、全量随机结果、同种子重跑；计数和有限模拟分辨率同时报告。']],[3.0,13.3])
h('3.2 Skill与裸Prompt对比',2)
table(['维度','裸Prompt／原型','修订Skill'],[['检验','R2：300次、种子825200520','相同数据、模型、分配规则和种子'],['数值','0次超越，直接比例0/300','仍为0次；300个系数最大差0'],['解释','没有量化模拟分辨率，易被误读成p=0','加一平滑1/301=0.003322；单侧95%上界0.009936；限定为模拟诊断'],['可审计性','保留回归及种子','增加契约、吸收项、前后哈希和完整状态']],[2.2,6.5,7.6])
p('差异的本质：把需重复提醒的验证条件写入程序和制度，不以更显著的数字证明Skill有效。裸Prompt已有完整随机分布，不虚构“原版什么都没有”的对照。')
h('3.3 一般适用性自检',2)
p('迁移到重新生成的最低工资与就业合成面板时，重新确定地区、年份、对数就业、政策暴露及外生控制，改变采用时点和收益方向；Skill源码不改。80地区、8年、120次模拟完成；重复键、处理反转与交错采用均触发停止。通用的是检查制度，同步平衡面板辅助程序有明确适用范围，不能推广为任意DID都可自动运行。')
page();h('四、我设计的Agent')
h('4.1 结构摘要',2)
table(['要素','我的设计'],[['名称','eval-did-robustness'],['目标','把主回归与全部诊断转化为有证据、有限定的综合判断。'],['输入','识别目标、变量映射、主结果与参考实现、各项CSV、事件结果、安慰剂重复、诊断和汇总。'],['S1—S2','读取证据并核对口径；按通过、警示、失败逐项判定。通过只表示局部诊断未发现预定问题。'],['S3—S4','分别评价方向、量级、精度及估计对象；列出趋势、溢出、测量、采用选择、小簇和外部效度威胁。'],['S5—S6','解释能力互补性与异质性，再列可以写、须限定后写、不能写的具体语句。'],['输出','主结果、判定表、稳定性、残余威胁、主题解读、写作建议、边界声明及证据来源。'],['边界','工具仅Read/Glob/Grep；不跑回归、不改论文。调用者保存报告并验证受保护文件哈希。'],['验证','先校验必备文件、规格覆盖、数字范围和跨CSV一致性；在临时副本进行错误输入测试。']],[3.0,13.3])
h('4.2 设计心得',2)
p('最容易遗漏的是：Agent说明写了“缺失就停”，配套程序却未实际检查。初版报告已经包含主题解释，但输入门槛不足。第二轮不是重写一遍结论，而是用缺失与篡改样例验证停止行为，再修复缺口。')
p('评估最有价值的部分是把三种现象分开：R5仍显著但精度不稳定；R7下降涉及动态平均对象改变；T3未获支持属于研究假说证据不足，不是程序失败。这些区别直接约束论文措辞。')
p('实现说明：本次由Codex按照Agent文件完成语义评估，evaluate.py进行确定性校验和装配；未声称调用独立外部模型。通用Agent说明与课程项目适配脚本分开，迁移时需重新配置输入及主题内容。')
page();h('五、迭代反思')
h('第一轮  Skill的有限模拟解释',2)
p('观察问题：R2裸Prompt与Skill原型都得到tail_count=0、tail_rate=0。零次超越是计数事实，但不能推出尾部概率为零。原型虽保留全量抽样，未量化有限重复的不确定性。')
p('修订：在Skill验证清单与辅助程序加入加一平滑比例、模拟分辨率、零超越上界、失败重复状态和原数据哈希。固定数据、种子与模型重跑，避免把结果差异误当设计改进。')
p('修订前输出：tail_count: 0；tail_rate: 0.0。')
p('修订后输出：smoothed_tail_rate: 0.003322259；tail_probability_upper95: 0.009936082；status: complete。300个系数最大绝对差为0。')
p('证据：4e57924（原型）→819ad98（修订）；output/iterations/skill_v1、output/skill_r2、output/skill_comparison.json。',True)
h('第二轮  Agent的输入验证',2)
p('观察问题：在临时副本删除事件文件、改基准、删短窗口规格、令p=1.2，初版校验仍接受；只有重复行被识别。六例含一个正常样本，初版仅2/6符合预期。这里没有捏造AI先前说过的话，证据是实际程序行为。')
p('修订：增加必备输入与规格契约、有限数与概率范围、分项结果交叉核对、安慰剂重复校验及哈希检查。正常输入继续接受，五类错误全部停止。')
p('修订前输出：missing_event_file → accepted: true；changed_baseline → accepted: true。')
p('修订后输出：missing_event_file → STOP: missing required evidence；changed_baseline → STOP: source mismatch baseline。6/6符合预期。')
p('证据：8b216e6（初版）→9049b5c（修订）；output/iterations/agent_v1_tests.json、agent_v2_tests.json。',True)
h('迭代总结',2)
p('需要改变的不是让AI说得更谨慎，而是让关键判断有可执行的依据。下一次研究应先定义估计对象和异常退出条件，再运行模型；保留不利结果、实际失败和版本差异，才能解释输出为何值得相信。仍需人工复核经济语义与研究假设，测试通过不是因果证明。')
page();h('六、Git与协作习惯')
table(['问题','我的回答'],[['分支设置','main之外5个任务分支，分别对应稳健性、Skill、Agent、论文和迭代。每个任务先建分支，完成后合并。'],['diff审查','合并前审查差异和关键代码；修正课堂公式冗余项，T4改为显式基准交互，保留异常输出。日志的尾部空格属于编译器格式，另行处理。'],['commit规范','feat / fix / docs / chore / merge，说明动作与原因。代表提交：9049b5c，增加证据完整性校验并拒绝被破坏的输入。'],['Git的作用','把实际尝试、拒绝的输出、修复和最终证据对应到版本；已有历史来自本次工作，不复制同学历史，不补造旧日期。']],[3.0,13.3])
h('Git历史文本快照',2)
history=(ROOT/'output/git-history.txt').read_text() if (ROOT/'output/git-history.txt').exists() else subprocess.check_output(['git','log','--oneline','--graph','--all'],cwd=ROOT,text=True)
for line in history.splitlines():
 pp=doc.add_paragraph(line);pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1
 for run in pp.runs:run.font.name='Menlo';run.font.size=Pt(7.3)
md.append('```text\n'+history+'```\n')
p('快照保存于生成材料时；后续提交以仓库 git log --oneline --graph --all 为准。远程仓库由本人最后创建并推送。',True)
page();h('附录  提交物自查清单')
table(['提交物','位置','状态'],[['完整Git仓库','did-homework/.git','已完成'],['识别目标','output/estimand.md','已完成'],['Prompt','prompts/robustness/','基准+12张'],['检验汇总','output/robustness_summary.md','已完成'],['全部结果','output/*.csv、*.png','已完成'],['Skill','.claude/skills/robustness-check/','已验证'],['Skill对照','audit-log.md、output/skill_comparison.json','已完成'],['Agent','.claude/agents/eval-did-robustness.md','已完成'],['综合评估','output/eval_robustness_report.md','已完成'],['论文与PDF','paper/main.tex、main.pdf','实际编译'],['AI使用记录','audit-log.md','已完成'],['两轮迭代','output/iteration_reflection.md','已完成'],['Git历史','output/git-history.txt','已完成'],['纸质材料','submission/纸质提交材料.docx及PDF','已生成'],['远程推送与链接','按GITHUB_UPLOAD.md操作','待本人完成']],[3.3,10.8,2.2])
h('关键证据速查',2)
p('复现：output/baseline_comparison.md。推断口径：diagnostics.json及all_results.csv。原始公式异常：output/iterations/baseline_original_script.csv。迁移：output/migration/README.md和negative_tests.json。文献核查：quality_reports/crossref_verified.json。')
p('论文与纸质材料均保留合成数据边界。材料中没有虚构现实数据、参考文献、独立模型调用或GitHub地址；提交前需将本人真实仓库链接填入封面，确认内容后打印。')
# Remove inherited template borders and theme font overrides; set explicit CJK runs.
for element in doc.styles.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
for element in doc._element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
for fonts in doc.styles.element.xpath('.//w:rFonts'):
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+attr),None)
    fonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
for root in [doc._element,sec.footer._element]:
    for rr in root.xpath('.//w:r'):
        pr=rr.find(qn('w:rPr'))
        if pr is None:pr=OxmlElement('w:rPr');rr.insert(0,pr)
        fonts=pr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');pr.insert(0,fonts)
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+attr),None)
        fonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        if not fonts.get(qn('w:ascii')):
            fonts.set(qn('w:ascii'),'Noto Serif CJK SC');fonts.set(qn('w:hAnsi'),'Noto Serif CJK SC')
doc.core_properties.author='邓鸿琪' ;doc.core_properties.title='AI时代经济学研究工具与方法课程作业纸质提交材料';doc.save(OUT/'纸质提交材料.docx');(OUT/'纸质提交材料.md').write_text('\n'.join(md),encoding='utf-8')
print(OUT/'纸质提交材料.docx')
